"""Verify editable drivers using disposable Artifact Tool recalculations."""
import argparse
import hashlib
import json
import os
import pathlib
import shutil
import subprocess
import sys
import tempfile

import openpyxl

from verify_model import package


def main():
    directory = pathlib.Path(__file__).resolve().parent
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", nargs="?", default=str(directory / "cost-model.xlsx"))
    parser.add_argument("--report")
    args = parser.parse_args()
    source = pathlib.Path(args.workbook).resolve()
    before_hash = hashlib.sha256(source.read_bytes()).hexdigest()
    baseline = openpyxl.load_workbook(source, data_only=True)
    addresses = {row[0].value: row[1].coordinate for row in baseline["Assumptions"].iter_rows(min_row=5, max_col=2) if row[0].value is not None}
    value = lambda workbook, sheet, address: float(workbook[sheet][address].value)
    node = os.environ.get("ARTIFACT_TOOL_NODE") or shutil.which("node")
    assert node, "Set ARTIFACT_TOOL_NODE to the Artifact Tool runtime Node executable"
    scenarios = [
        ("enrolment", "Wave 1 employees (RD + PK, Thailand)", value(baseline, "Assumptions", addresses["Wave 1 employees (RD + PK, Thailand)"]) + 1000),
        ("subscription_price", "Subscription price per enrolled employee per month", value(baseline, "Assumptions", addresses["Subscription price per enrolled employee per month"]) + 4),
        ("ai_usage_double", "Coach messages per active user per month", value(baseline, "Assumptions", addresses["Coach messages per active user per month"]) * 2),
        ("ai_usage_zero", "Coach messages per active user per month", 0),
    ]
    results = []
    with tempfile.TemporaryDirectory(prefix="iris-financial-verification-") as temporary:
        temporary = pathlib.Path(temporary)
        for name, label, changed in scenarios:
            output = temporary / f"{name}.xlsx"
            subprocess.run([node, str(directory / "recalculate.mjs"), f"--input={source}", f"--output={output}", f"--set=Assumptions!{addresses[label]}={changed}"], check=True, capture_output=True, text=True)
            assert package(source)["formulas"] == package(output)["formulas"], f"{name}: formulas changed"
            subprocess.run([os.environ.get("ARTIFACT_TOOL_PYTHON") or sys.executable, str(directory / "verify_model.py"), str(output), f"--report={temporary / (name + '.json')}"], check=True, capture_output=True, text=True)
            candidate = openpyxl.load_workbook(output, data_only=True)
            rev_delta = value(candidate, "Pricing & P&L", "E8") - value(baseline, "Pricing & P&L", "E8")
            cost_delta = value(candidate, "Pricing & P&L", "E9") - value(baseline, "Pricing & P&L", "E9")
            profit_delta = value(candidate, "Pricing & P&L", "E10") - value(baseline, "Pricing & P&L", "E10")
            if name == "enrolment":
                enrolled_delta = changed - value(baseline, "Assumptions", addresses[label])
                months = sum(1 for month in range(1, 37) if month >= value(baseline, "Assumptions", addresses["Wave 1 go-live month"]))
                expected = enrolled_delta * months * value(baseline, "Assumptions", addresses["Subscription price per enrolled employee per month"])
                assert abs(rev_delta - expected) <= 1 and cost_delta > 0
                assert value(candidate, "Run", "E5") - value(baseline, "Run", "E5") == enrolled_delta
            elif name == "subscription_price":
                billed = sum(float(cell.value) for cell in baseline["Pricing & P&L"]["B32:AK32"][0])
                expected = billed * (changed - value(baseline, "Assumptions", addresses[label]))
                assert abs(rev_delta - expected) <= 1 and abs(cost_delta) <= 1 and profit_delta > 0
            else:
                assert abs(rev_delta) <= 1
                assert (cost_delta > 0 and profit_delta < 0) if name.endswith("double") else (cost_delta < 0 and profit_delta > 0)
                if name.endswith("zero"):
                    assert value(candidate, "Run", "AK8") == 0 and value(candidate, "Run", "AK9") == 0
                    assert value(candidate, "Run", "AK13") > 0, "Path generation/evaluation costs remain at zero coach usage"
            results.append({"scenario": name, "input": f"Assumptions!{addresses[label]}", "new_value": changed, "revenue_delta_thb": rev_delta, "cost_delta_thb": cost_delta, "gross_profit_delta_thb": profit_delta, "status": "PASS"})
    assert before_hash == hashlib.sha256(source.read_bytes()).hexdigest(), "Sensitivity tests altered the deliverable"
    report = {"status": "PASS", "source_unchanged": True, "scenarios": results, "engine": "Artifact Tool public import/recalculate/export API", "native_excel": "SKIPPED: no desktop Excel session was used"}
    if args.report:
        pathlib.Path(args.report).write_text(json.dumps(report, indent=2) + "\n")
    print(json.dumps(report, indent=2))


if __name__ == "__main__":
    main()
